"""Landscape visual report and product catalogue from the archived public audit."""
import json, math
from pathlib import Path
from urllib.parse import urlsplit
from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'docs/diocese-catalog'; IMG=OUT/'images'
VIS=IMG/'visual-report';VIS.mkdir(exist_ok=True)
DATA=json.loads((OUT/'research/inventory.json').read_text(encoding='utf-8'))
MOBILE=json.loads((OUT/'research/mobile.json').read_text(encoding='utf-8'))
FONT='C:/Windows/Fonts/arial.ttf'; BOLD='C:/Windows/Fonts/arialbd.ttf'
INK='#20374A';GOLD='#A87927';RED='#B24E42';GREEN='#39766E';GREY='#E4E8EA'

def font(size=30,bold=False):return ImageFont.truetype(BOLD if bold else FONT,size)
def wrap(draw,text,width,size=28,bold=False):
    lines=[];line=''
    for word in text.split():
        test=(line+' '+word).strip()
        if draw.textlength(test,font=font(size,bold))>width and line:lines.append(line);line=word
        else:line=test
    if line:lines.append(line)
    return '\n'.join(lines)

def gallery(name,items,height=840):
    canvas=Image.new('RGB',(1800,height),'white');dr=ImageDraw.Draw(canvas)
    cell=1800//len(items)
    for i,(path,title,note) in enumerate(items):
        x=i*cell+12
        dr.multiline_text((x,5),wrap(dr,title,cell-28,30,True),font=font(30,True),fill=INK,spacing=4)
        im=Image.open(path).convert('RGB')
        thumb=ImageOps.contain(im,(cell-34,height-165))
        canvas.paste(thumb,(x+(cell-28-thumb.width)//2,88))
        dr.multiline_text((x,height-65),wrap(dr,note,cell-28,22),font=font(22),fill=INK,spacing=3)
    target=VIS/(name+'.png');canvas.save(target);return target

def bars(name,rows):
    canvas=Image.new('RGB',(1800,700),'white');dr=ImageDraw.Draw(canvas)
    for i,(title,a,b,left,right) in enumerate(rows):
        y=25+i*220;dr.text((30,y),title,font=font(34,True),fill=INK)
        total=a+b;length=1700;cut=int(length*a/total)
        dr.rectangle((30,y+66,30+cut,y+126),fill=GREEN)
        dr.rectangle((30+cut,y+66,30+length,y+126),fill=GOLD)
        dr.text((30,y+143),f'{a} из {total}  ·  {a/total:.1%}  {left}'.replace('.',','),font=font(28),fill=GREEN)
        dr.text((930,y+143),f'{b} из {total}  ·  {b/total:.1%}  {right}'.replace('.',','),font=font(28),fill=GOLD)
    target=VIS/(name+'.png');canvas.save(target);return target

def errors():
    im=Image.new('RGB',(1800,680),'white');dr=ImageDraw.Draw(im)
    for i,(number,title,text) in enumerate([(3,'Имя сайта не найдено','Пфорцхайм · Вупперталь · община Кирилла и Мефодия в Гамбурге'),(1,'Ошибка сертификата','Зулинген · имя в сертификате не совпадает с адресом'),(1,'Требуется авторизация','Бад Кройцнах · HTTP 401')]):
        x=i*600+25;dr.text((x,20),str(number),font=font(150,True),fill=RED)
        dr.multiline_text((x,210),wrap(dr,title,550,34,True),font=font(34,True),fill=INK,spacing=8)
        dr.multiline_text((x,335),wrap(dr,text,540,30),font=font(30),fill=INK,spacing=10)
        dr.text((x,590),'Подтверждено повторным запросом',font=font(23),fill=RED)
    p=VIS/'errors.png';im.save(p);return p

def document():
    d=Document();s=d.sections[0];s.orientation=WD_ORIENT.LANDSCAPE
    s.page_width=Inches(11);s.page_height=Inches(8.5)
    s.top_margin=Inches(.58);s.bottom_margin=Inches(.55);s.left_margin=s.right_margin=Inches(.65)
    for name,size in [('Normal',11),('Title',25),('Subtitle',13),('Heading 1',21),('Heading 2',14),('Caption',9)]:
        st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.font.bold=name.startswith('Heading');st.paragraph_format.space_after=Pt(7)
        for b in st.element.xpath('.//w:pBdr'):b.getparent().remove(b)
    d.styles['Normal'].paragraph_format.line_spacing=1.06
    s.header.paragraphs[0].text='ПРИХОД  |  Материалы для Владыки Тихона  |  15 октября 2026'
    s.header.paragraphs[0].style='Caption'
    p=s.footer.paragraphs[0];p.alignment=2;p.add_run('Приход  ·  ')
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
    d.core_properties.author='Проект Приход';return d

def page(d,title,text='',image=None,caption='',first=False):
    if not first:d.add_page_break()
    d.add_heading(title,0 if first else 1)
    if text:d.add_paragraph(text)
    if image:
        im=Image.open(image);w,h=im.size;display=min(9.6,4.6*w/h)
        p=d.add_paragraph();p.alignment=1;p.add_run().add_picture(str(image),width=Inches(display))
    if caption:d.add_paragraph(caption,style='Caption')

def shot(n):return IMG/f'mobile-{n:02}.png'
def g(name,items):return gallery(name,items)

def fragment(name,path,box,highlight=None):
    im=Image.open(path).convert('RGB').crop(box)
    if highlight:ImageDraw.Draw(im).rectangle(highlight,outline=RED,width=4)
    target=VIS/(name+'.png');im.save(target);return target

def comparison(name,left,right,left_note,right_note):
    im=Image.new('RGB',(1800,840),'white');dr=ImageDraw.Draw(im)
    for x,path,label,color,note in [(20,left,'ПРОБЛЕМА',RED,left_note),(980,right,'РЕШЕНИЕ В ПРИХОДЕ',GREEN,right_note)]:
        dr.text((x,5),label,font=font(39,True),fill=color)
        thumb=ImageOps.contain(Image.open(path).convert('RGB'),(780,610))
        im.paste(thumb,(x+(800-thumb.width)//2,85))
        dr.multiline_text((x,725),wrap(dr,note,790,32,True),font=font(32,True),fill=color,spacing=8)
    dr.line((848,370,948,370),fill=INK,width=8)
    dr.polygon([(948,370),(918,350),(918,390)],fill=INK)
    target=VIS/(name+'.png');im.save(target);return target
def summary(d):
    page(d,'Обзор сайтов епархии',
         'Проверены все 102 карточки из официального списка. Адреса относятся к 68 доменам; один сайт может обслуживать несколько общин.',
         bars('summary',[('Есть ли ссылка в карточке прихода',76,26,'ссылка есть','ссылка не указана'),('Открывается ли опубликованный адрес',65,5,'получен HTML','повторная ошибка'),('Помещается ли страница в экран 390 px',45,20,'нет превышения ширины','ширина превышена')]),
         'Срез 15.10.2026 · rokmp.de/churches-by-cities/. Пять ошибок: 3 DNS, 1 сертификат, 1 запрос авторизации. Отсутствие ссылки не доказывает отсутствия сайта; HTML ответ не гарантирует его работу. Посещаемость не измерена. Полный реестр и все 65 снимков сохранены отдельно.',first=True)

def overview():
    d=document();summary(d)
    cropped=fragment('compact-mobile-problem',shot(15),(0,210,375,780))
    cards=fragment('compact-cards',IMG/'project-schedule-fragment.png',(15,270,365,885))
    image=comparison('compact-mobile',cropped,cards,'Строки уходят за край экрана','Время и название службы видны')
    page(d,'Человек должен прочитать расписание с телефона',
         'Нюрнберг: даты актуальны, но строки обрезаны на экране 390 px. Решение — готовые карточки, которые помещаются по ширине телефона.',image,
         'Что делает приход: выбирает оформление и ведёт службы через готовые поля. Слева: фрагмент rus-kirche.de, 15.10.2026. Справа: реальный демонстрационный интерфейс «Прихода»; чужой сайт не изменён.')
    old=fragment('compact-old-dates',shot(32),(125,155,365,550),(5,65,85,350))
    editor=fragment('compact-editor',ROOT/'docs/design-reference/verification-2026-10-12/local-calendar-editor-rc10.jpg',(220,500,785,990))
    image=comparison('compact-dates',old,editor,'В ближайших службах даты 2025 года','Сотрудник вводит дату и время')
    page(d,'Обновить расписание должно быть просто',
         'Тюбинген: в разделе ближайших служб остались даты 16–30 августа 2025 года. В «Приходе» дата, время и место задаются обычными полями.',image,
         'Что делает приход: добавляет службу, проверяет дату и публикует; регулярные службы настраивает повторением. Автоматическая актуальность не обещается. Слева: orthodoxekirche.de/Tuebingen/, срез 15.10.2026; справа: фрагмент проверочного редактора проекта.')
    picture=fragment('compact-image-schedule',shot(18),(0,80,375,630))
    image=comparison('compact-structured',picture,cards,'Расписание картинкой нужно увеличивать','Крупные отдельные карточки служб')
    page(d,'Расписание читается без увеличения картинки',
         'Мюнхен: на первом экране виден лист расписания с мелким текстом. В «Приходе» каждая служба представлена отдельной карточкой с временем и названием.',image,
         'Что делает приход: вносит службы в редактор, а готовый блок показывает их на сайте; не нужно заново оформлять изображение расписания для каждого изменения. Слева: voskresenie.de, 15.10.2026. Справа: фрагмент демонстрации проекта.')
    frames=fragment('compact-frames-problem',IMG/'brandenburg-mobile.png',(0,180,375,760))
    home=fragment('compact-home',ROOT/'docs/test-server/monastery-home-mobile.png',(0,220,375,1000))
    image=comparison('compact-navigation',frames,home,'Меню зажато между рамками','Одна страница и заметная кнопка')
    page(d,'Посетителю нужен понятный путь к богослужениям',
         'Бранденбург: меню и содержимое живут в отдельных узких рамках. Готовая мобильная страница «Прихода» даёт понятную кнопку расписания и ближайшую службу.',image,
         'Что делает приход: выбирает готовое оформление и заполняет сведения; структуру мобильной страницы собирает тема. Слева: brandenburg.orthodox.bplaced.de/ru/index.html, 15.10.2026. Справа: фрагмент действующей демонстрации, не макет обновления чужого сайта.')
    deleted=fragment('compact-deleted',shot(45),(0,0,375,110))
    image=comparison('compact-start',deleted,IMG/'project-classic.png','По ссылке сообщение об удалении','Готовая основа нового сайта')
    page(d,'Приход может сам наполнить готовый сайт',
         'Саарбрюккен: опубликованный адрес сообщает об удалении сайта. После проверки актуального адреса общине при необходимости можно предложить запуск «Прихода».',image,
         'Как исправить: технический сопровождающий проверяет домен, HTTPS и установку; приход проходит мастер и заполняет сайт. Мы готовы помочь. Это предложение, а не выполненный переезд. Все 65 снимков и полный реестр сохранены отдельно; отсутствие ссылки не доказывает отсутствия сайта.')
    d.save(OUT/'Diocese_sites_appendix.docx')

def proposal():
    d=document()
    page(d,'Проект Приход','Для Владыки Тихона, архиепископа Рузского, управляющего Берлинско Германской епархией. Предлагаем готовые средства, чтобы приход мог сам создать и вести полезный сайт. Мы готовы помочь при запуске и освоении.',IMG/'project-classic.png','Работающий демонстрационный продукт на материалах монастыря. Это иллюстрация возможностей, а не новый официальный сайт монастыря.',True)
    summary_image=bars('proposal-stats',[('Ссылка на сайт в карточке',76,26,'указана','не указана'),('Результат открытия адреса',65,5,'получен HTML','повторная ошибка'),('Мобильная ширина страницы',45,20,'без превышения','ширина превышена')])
    page(d,'Почему тема заслуживает внимания','Проверка епархиального каталога показывает разные потребности: уточнение адресов, обновление расписания и улучшение чтения с телефона. В коротком обзоре показаны конкретные проблемы и решения; полный набор снимков сохранён отдельно.',summary_image,'102 карточки · 70 адресов на 68 доменах · 65 мобильных просмотров · данные на 15.10.2026. Отсутствие ссылки не доказывает отсутствия сайта; посещаемость не установлена.')
    page(d,'Сотрудник работает с готовыми средствами','После технической установки сотрудник проходит мастер, заполняет сведения, выбирает страницы и оформление, проверяет план. Публикация выполняется отдельно после подтверждения.',ROOT/'docs/design-reference/verification-2026-10-12/local-wizard-plan-rc9.jpg','Реальный интерфейс мастера из проверочного набора 12.10.2026. Мастер предусматривает повторный запуск без дублей. Обычное наполнение не требует редактирования кода.')
    for key,title in [('classic','Классическое оформление'),('hospitality','Светлое оформление'),('editorial','Журнальное оформление'),('tradition','Традиционное оформление'),('blue','Современное оформление')]:
        page(d,title,'Приход использует собственные название, фотографии, контакты и материалы. Готовая структура позволяет выбрать внешний вид без отдельного проектирования каждого раздела.',IMG/f'project-{key}.png','Снимок демонстрационного сайта из ресурсов темы. Оформление не заменяет актуального содержания; за сведения отвечает приход.')
    page(d,'Расписание и обращения в повседневной работе','Сотрудник ведёт регулярные и разовые службы, переносы и отмены. Посетитель видит ближайшую службу и может задать вопрос без регистрации. Ответственные работают с обращениями в закрытом кабинете.',g('proposal-daily',[(IMG/'project-schedule-fragment.png','Расписание','Неделя месяц список и печать'),(IMG/'project-question-fragment.png','Вопрос священнику','Назначение ответственного и ответ')]),'Календарь даёт сведения о праздниках, но фактические службы определяет приход. Для обращений необходимо назначить отвечающего и проверить получение почты.')
    page(d,'Готовность и условия применения','Основной продукт реализован в WordPress и имеет статус release candidate. Профиль, места служения, духовенство, расписание, обращения, права и мастер входят в основной продукт.')
    for title,text in [('Четыре языка интерфейса','Русский, немецкий, польский и украинский. Содержание прихода не переводится автоматически; языковые версии нужно подготовить.'),('Необязательные подключения','Календарь, записки, пожертвования и внешние сервисы подключаются отдельно. Оплата и доставка сообщений требуют настройки и проверки; WhatsApp и Facebook не выдаются за полностью проверенные живые интеграции.'),('Техническое сопровождение','Нужны хостинг, домен, HTTPS, почта, обновления и резервные копии. Переезд с произвольного конструктора не заявляется автоматическим.'),('Ответственность за содержание','Приход определяет расписание, публикует материалы и отвечает людям. Единой готовой панели епархии и статистики посещаемости в продукте не заявляется.')]:d.add_heading(title,2);d.add_paragraph(text)
    page(d,'Предлагаем проверить самостоятельную работу в пилоте','Просим благословить добровольное опытное применение в трёх заинтересованных приходах разного размера. Число участников и предлагаемые четыре недели наблюдения требуют согласования.')
    for title,text in [('Подготовка и запуск','Согласовать участников, ответственных, условия помощи и резервную копию. Установить продукт, вместе пройти мастер и заполнить первые сведения.'),('Самостоятельные задачи','Сотрудник сам меняет время службы, оформляет отмену, обновляет контакты, публикует новость и обрабатывает тестовое обращение. Записываются время, ошибки и случаи помощи.'),('Решение по результату','Оценить актуальность расписания, удобство для посетителей, самостоятельность сотрудника и реальные расходы. Сроки, стоимость и объём поддержки согласуются заранее; измеренной экономии пока нет.'),('Просьба Владыке','Благословить демонстрацию заинтересованным настоятелям и определить контактное лицо для согласования пилота. Команда готова помочь с запуском и обучением в согласованном объёме.')]:d.add_heading(title,2);d.add_paragraph(text)
    d.save(OUT/'Prihod_proposal.docx')

def brief():
    d=document()
    page(d,'Приход может вести сайт самостоятельно','Ваше Высокопреосвященство, Владыка Тихон! Предлагаем проект «Приход» и просим рассмотреть добровольное опытное применение в нескольких приходах. Мы даём готовые средства и готовы помочь с техническим запуском и обучением.',IMG/'project-classic.png','Сотрудник проходит мастер, выбирает оформление и работает с готовыми формами. Обычное обновление расписания, контактов и материалов не требует редактирования программного кода.',True)
    page(d,'Практическое предложение','Предлагаем пилот в трёх приходах с четырьмя неделями наблюдения после запуска. Участники, сроки, стоимость и объём помощи согласуются до начала.')
    for title,text in [('Основание','Проверены 102 карточки, 70 опубликованных адресов и 65 мобильных страниц. В пяти случаях повторилась ошибка открытия, ещё один адрес сообщает об удалении сайта. В 20 мобильных просмотрах ширина документа превышает экран. Подробности — в иллюстрированном обзоре.'),('Результат для прихода','Сотрудник сам ведёт службы, переносы и отмены, обновляет контакты, публикует новости и отвечает на вопросы. Мы фиксируем где нужна помощь и совершенствуем продукт по реальным замечаниям.'),('Условия','WordPress и техническое обслуживание требуют ответственного. Основной продукт имеет статус release candidate; календарь, записки, пожертвования и внешние сервисы настраиваются отдельно. Тексты прихода не переводятся автоматически.'),('Просьба','Благословить демонстрацию заинтересованным настоятелям и согласование пилота через контактное лицо. Посещаемость без доступа к аналитике не установлена; отсутствие ссылки в карточке не означает отсутствия сайта.')]:d.add_heading(title,2);d.add_paragraph(text)
    d.add_paragraph('Срез 15.10.2026 · https://rokmp.de/churches-by-cities/ · Демонстрация https://prihod.georg-kloster.ru/',style='Caption')
    d.save(OUT/'Prihod_brief.docx')

if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser();parser.add_argument('--target',choices=['all','overview','proposal','brief'],default='all');args=parser.parse_args()
    for name,build in [('overview',overview),('proposal',proposal),('brief',brief)]:
        if args.target in ('all',name):build()
    print('Built requested landscape documents')
